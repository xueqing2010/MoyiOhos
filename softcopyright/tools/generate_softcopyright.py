#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate StarFlow software-copyright application materials.

Run with the bundled Python that has python-docx:
  /Users/kong/.dsh/dsh-runtimes/dsh-next-primary-runtime/dependencies/python/bin/python3       softcopyright/tools/generate_softcopyright.py

The generated DOCX files are editable. Convert them to PDF in Word/Pages
before uploading to the China Copyright Protection Center system.
"""
from __future__ import annotations

import json
import re
import textwrap
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "softcopyright"
IMAGES = ROOT / "docs" / "images"
SOFTWARE_NAME = "星译漫画小说智能翻译软件"
SOFTWARE_SHORT = "星译"
VERSION = "V1.1.0"
APPLICANT = "【申请人名称，待填】"
DATE_TEXT = "2026年10月"

SOURCE_EXTS = {".ets", ".cpp", ".h"}
EXCLUDE_PARTS = {
    ".git", "build", "oh_modules", "node_modules", ".cxx", ".preview",
    "third_party", "llama.cpp", "softcopyright", "entry/build",
}
SOURCE_WRAP = 96
PAGE_LINES = 50
SAMPLE_PAGES = 60


def set_cn_font(run, name: str, size: float | None = None, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold


def add_field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def page_setup(doc: Document, margins: float = 1.6) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(margins)
    section.bottom_margin = Cm(margins)
    section.left_margin = Cm(margins)
    section.right_margin = Cm(margins)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)


def add_header_footer(doc: Document, header_text: str) -> None:
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run(header_text)
    set_cn_font(run, "宋体", 8)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("第 ")
    set_cn_font(r1, "宋体", 8)
    add_field(fp, "PAGE")
    r2 = fp.add_run(" 页 / 共 ")
    set_cn_font(r2, "宋体", 8)
    add_field(fp, "NUMPAGES")
    r3 = fp.add_run(" 页")
    set_cn_font(r3, "宋体", 8)


def collect_sources() -> list[Path]:
    files: list[Path] = []
    for p in ROOT.rglob("*"):
        if not p.is_file() or p.suffix not in SOURCE_EXTS:
            continue
        if p.name == "BuildProfile.ets":
            continue
        rel = p.relative_to(ROOT)
        if any(part in EXCLUDE_PARTS for part in rel.parts):
            continue
        if str(rel).startswith("entry/.cxx") or str(rel).startswith("entry/.preview"):
            continue
        files.append(p)
    def key(p: Path) -> tuple:
        rel = str(p.relative_to(ROOT))
        priority = 0
        if rel.startswith("entry/src/main/ets/entryability"):
            priority = 1
        elif rel.startswith("entry/src/main/ets/pages"):
            priority = 2
        elif rel.startswith("entry/src/main/ets/"):
            priority = 3
        elif rel.startswith("core"):
            priority = 4
        elif rel.endswith(".cpp") or rel.endswith(".h"):
            priority = 5
        return (priority, rel)
    return sorted(files, key=key)


def build_source_lines(files: list[Path]) -> tuple[list[str], int]:
    lines: list[str] = [
        "// 软件名称：" + SOFTWARE_NAME,
        "// 版本号：" + VERSION,
        "// 著作权人：" + APPLICANT,
        "// 本文件为计算机软件著作权登记源程序鉴别材料（前30页+后30页）。",
        "//",
    ]
    for p in files:
        rel = str(p.relative_to(ROOT))
        lines.append("// ==================== 文件：" + rel + " ====================")
        try:
            raw = p.read_text(encoding="utf-8", errors="ignore")
        except OSError:
            continue
        for raw_line in raw.splitlines():
            line = raw_line.replace("\t", "    ").rstrip()
            if len(line) <= SOURCE_WRAP:
                lines.append(line)
            else:
                wrapped = textwrap.wrap(line, width=SOURCE_WRAP, replace_whitespace=False,
                                        drop_whitespace=False, break_long_words=True,
                                        break_on_hyphens=False)
                lines.extend(wrapped if wrapped else [""])
    total = len(lines)
    if total > SAMPLE_PAGES * PAGE_LINES:
        half = (SAMPLE_PAGES // 2) * PAGE_LINES
        lines = lines[:half] + lines[-half:]
    return lines, total


def write_source_docx(lines: list[str], out_docx: Path, out_txt: Path) -> None:
    out_txt.write_text("\n".join(lines), encoding="utf-8")
    doc = Document()
    page_setup(doc, 1.2)
    add_header_footer(doc, SOFTWARE_NAME + " " + VERSION + "  源程序鉴别材料")
    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(7.5)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    style.paragraph_format.line_spacing = Pt(9.5)
    total_pages = SAMPLE_PAGES
    for page in range(total_pages):
        start = page * PAGE_LINES
        chunk = lines[start:start + PAGE_LINES]
        for i in range(PAGE_LINES):
            text = chunk[i] if i < len(chunk) else ""
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(9.5)
            run = p.add_run(text if text else " ")
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(7.5)
            if text.startswith("//"):
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        if page < total_pages - 1:
            doc.add_page_break()
    props = doc.core_properties
    props.title = SOFTWARE_NAME + " " + VERSION + " 源程序鉴别材料"
    props.author = APPLICANT
    props.subject = "计算机软件著作权登记源程序"
    doc.save(out_docx)


def add_para(doc: Document, text: str, size: float = 10.5, bold: bool = False,
             first_indent: bool = True, align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_cn_font(run, "宋体", size, bold)


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_cn_font(run, "黑体", 15, True)


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_cn_font(run, "黑体", 12, True)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.4
    run = p.add_run(text)
    set_cn_font(run, "宋体", 10.5)


def add_image(doc: Document, filename: str, caption: str, width_cm: float = 11.5) -> None:
    path = IMAGES / filename
    if not path.exists():
        add_para(doc, "（截图待补充：" + filename + "）", 9, False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("图：" + caption)
    set_cn_font(r, "宋体", 8.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def write_manual(out_docx: Path) -> None:
    doc = Document()
    page_setup(doc, 2.0)
    add_header_footer(doc, SOFTWARE_NAME + " " + VERSION + "  软件说明书")
    # cover
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(SOFTWARE_NAME)
    set_cn_font(r, "黑体", 26, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("软件说明书")
    set_cn_font(r, "黑体", 20, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("（计算机软件著作权登记文档鉴别材料）")
    set_cn_font(r, "宋体", 12)
    for _ in range(6):
        doc.add_paragraph()
    for label, value in [("软件版本", VERSION), ("著作权人", APPLICANT), ("编制日期", DATE_TEXT)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label + "：" + value)
        set_cn_font(r, "宋体", 11)
    doc.add_page_break()
    # toc
    add_h1(doc, "目录")
    toc = [
        "1 软件概述", "2 运行环境", "3 安装与启动", "4 主要功能",
        "5 操作流程", "6 多设备与沉浸式界面", "7 本地模型与性能",
        "8 数据与隐私", "9 常见问题与故障处理", "10 版权与合规说明",
    ]
    for item in toc:
        add_para(doc, item, 10.5, False)
    doc.add_page_break()
    # chapters
    chapters = [
        ("1 软件概述", [
            "星译漫画小说智能翻译软件是面向 HarmonyOS NEXT 的本地漫画、小说与文本翻译工具。软件支持导入本地漫画图片和 EPUB/TXT 小说，使用本地 OCR 识别漫画文字，结合在线翻译服务或本地 llama.cpp 模型生成译文，并把译文按原文版式覆盖回页面。",
            "软件采用 Stage 模型与多 HAR 模块架构：entry 负责页面与交互，coremodel 负责数据模型，coreocr 负责文字识别，coretranslate 负责翻译引擎，corerender 负责排版绘制，corenovel 负责小说解析。",
            "软件的主要目标是把漫画翻译流程从多步手工操作收敛为：导入、阅读、翻译、校对、导出。",
        ]),
        ("2 运行环境", [
            "硬件环境：HarmonyOS NEXT 手机、平板或 2in1 设备；建议内存 4GB 以上；本地 OCR 与本地翻译模型建议 8GB 以上。",
            "软件环境：HarmonyOS NEXT API 26 及以上；系统需具备图片解码、文件选择、网络访问能力。",
            "开发环境：DevEco Studio 6、ArkTS/ArkUI、C++ NAPI、CMake。",
            "网络环境：使用在线翻译时需要网络；本地 OCR 与本地翻译可离线运行。",
        ]),
        ("3 安装与启动", [
            "通过应用商店或调试安装 HAP 后，桌面出现“星译 StarFlow”图标。首次启动进入书架页，底部为书架、翻译、历史、我的四个入口。",
            "首次使用漫画翻译前，需要在“设置 → OCR”下载本地 PP-OCRv6 检测与识别模型；使用本地翻译时，需要在“模型管理”导入 GGUF 翻译模型。",
            "软件不强制注册账号，书籍、阅读进度、术语表、翻译记忆和偏好保存在本机。",
        ]),
        ("4 主要功能", [
            "漫画书架：导入本地漫画压缩包或图片目录，按漫画分组展示，支持最近阅读、未翻译筛选和整本预翻译入口。",
            "漫画阅读：支持左到右、右到左、竖排整页、连续 Webtoon 四种阅读流，支持翻页动画、背景色、反色、灰度、亮度、对比度和自动翻页。",
            "OCR 与翻译：本地 PP-OCRv6 检测文本行，原生 ONNX Runtime 识别，按空间位置合并气泡，调用在线或本地翻译引擎，按页面坐标排版覆盖。",
            "气泡校对：长按单行气泡打开校对框，修改译文后立即保存并重绘。",
            "小说翻译：导入 EPUB/TXT，按章节解析、分页阅读并翻译。",
            "导出：支持原文/译文/双语 ZIP，以及译文 PDF；导出文件保存在应用文件目录或通过系统保存选择器保存。",
            "术语表与翻译记忆：支持导入 JSON 术语表，翻译前替换专名；页翻译完成后自动积累翻译记忆。",
            "多设备布局：手机使用底部标签栏，平板/2in1 使用左侧导航轨，支持沉浸光感材质和宽屏双页。",
        ]),
        ("5 操作流程", [
            "5.1 导入漫画：在书架页点击“导入”，选择漫画压缩包或图片目录；导入完成后在书架卡片进入阅读。",
            "5.2 翻译本页：在阅读页点击页面中央打开菜单，选择“翻译本页”；进度条依次显示 OCR、翻译、保存阶段。",
            "5.3 校对：长按需要修改的单行气泡，在“校对译文”输入框中修改并保存；合并气泡暂不支持逐行编辑。",
            "5.4 导出：点击“导出ZIP”选择仅原文、仅译文、双语对照；或点击“导出 PDF”生成译文 PDF。",
            '5.5 术语表：设置 → 翻译 → 术语表 → 导入 JSON，格式为 [{ "s": "原文", "t": "译文" }]。',
            "5.6 本地模型：设置 → OCR 下载检测/识别模型；模型管理导入 .gguf 翻译模型，导入后复制到应用沙箱。",
        ]),
        ("6 多设备与沉浸式界面", [
            "软件使用断点监听窗口宽度：小屏使用底部标签栏，大屏使用左侧导航轨；阅读器在 lg 断点下支持双页并排。",
            "沉浸光感使用 ArkUI ImmersiveMaterial；标签栏与导航轨关闭交互光效，避免阅读时干扰。",
            "阅读器背景、亮度、字体大小、行距、横竖排和阅读方向均可独立设置，并保存在本机。",
        ]),
        ("7 本地模型与性能", [
            "本地 OCR：PP-OCRv6 det + rec，det 长边限制 768；原生 C++ ONNX Runtime 在 TaskPool 中运行 det/rec，主线程只解码 RGBA。",
            "本地翻译：llama.cpp + GGUF 模型，按空间位置合并 3–4 行气泡批量生成，减少重复推理。",
            "性能记录：模拟器 6 行合成页原生 OCR 约 594ms、翻译约 1.4s、合计约 2.1s；真机历史数据约 4.7s/页（具体随设备与模型变化）。",
            "模型文件不随安装包分发；用户自行下载或导入，模型许可由模型发布方决定。",
        ]),
        ("8 数据与隐私", [
            "书籍、页图、翻译结果、术语表、翻译记忆、阅读进度和 API Key 默认保存在本机应用沙箱。",
            "使用在线翻译时，待翻译文本会发送到用户选择的翻译服务商；使用 OpenAI 等境外服务可能涉及数据出境，用户应在设置中确认服务商和隐私政策。",
            "使用本地 OCR/本地翻译时，图片和文本不离开设备。",
            "卸载应用会删除应用沙箱数据；用户可在系统设置中清除应用数据。",
        ]),
        ("9 常见问题与故障处理", [
            "图片加载失败：检查文件是否被移动或删除；重新导入漫画。",
            "OCR 无结果：确认已下载 OCR 模型；复杂背景/超大框会被过滤；可尝试提高图片清晰度。",
            "翻译失败：检查网络、API Key、模型路径；在线服务可能限流或返回空。",
            "PDF 导出失败：确认设备支持 HMS PDFKit，且应用有文件写入权限；可改用 ZIP 导出。",
            "模型导入失败：确认文件为 .gguf 且完整；导入时会复制到沙箱，需预留存储空间。",
        ]),
        ("10 版权与合规说明", [
            "本软件是本地翻译工具，不内置漫画、小说或图片内容；用户应仅导入其有权使用的内容。",
            "软件包含第三方开源组件，具体见 NOTICE.md 与 softcopyright/03-第三方组件与模型清单.md。",
            "AI 翻译结果可能存在错误，用户应在导出或发布前自行校对；公开传播 AI 生成内容时应遵守生成合成内容标识等规定。",
            "软件名称、版本号、著作权人应与申请表、源程序页眉、说明书封面保持一致。",
        ]),
    ]
    for title, paras in chapters:
        add_h1(doc, title)
        for text in paras:
            add_para(doc, text)
        if title.startswith("3"):
            add_image(doc, "phone-shelf.jpg", "手机书架与导入入口", 9.5)
        elif title.startswith("4"):
            add_image(doc, "reader-original.jpg", "漫画阅读原文", 8.8)
            add_image(doc, "ocr-overlay.jpg", "译文覆盖层", 8.8)
        elif title.startswith("5"):
            add_image(doc, "text-translate.jpg", "文本翻译页面", 8.8)
        elif title.startswith("6"):
            add_image(doc, "tablet-wide.jpg", "平板宽屏与左侧导航轨", 13.0)
        elif title.startswith("7"):
            add_image(doc, "model-manager.jpg", "模型管理与存储占用", 8.8)
        doc.add_page_break()
    # appendix
    add_h1(doc, "附录 A 主要模块")
    for text in [
        "entry：应用入口、页面、服务、数据访问与设置。",
        "commons：主题、断点、日志、路径等公共能力。",
        "coremodel：漫画、小说、页面、翻译结果等数据模型。",
        "coreocr：系统 OCR、本地 PP-OCRv6 与原生 ONNX Runtime 桥。",
        "coretranslate：在线服务商、本地 llama.cpp、提示词与引擎注册。",
        "corerender：覆盖层排版、Canvas 绘制、字体与颜色。",
        "corenovel：EPUB/TXT 解析与章节模型。",
    ]:
        add_bullet(doc, text)
    add_h1(doc, "附录 B 版本记录")
    add_para(doc, "V1.1.0：新增原生 det+rec、Noto Sans CJK、PDF 导出、术语表、气泡校对；具体变更见 CHANGELOG.md。")
    props = doc.core_properties
    props.title = SOFTWARE_NAME + " " + VERSION + " 软件说明书"
    props.author = APPLICANT
    props.subject = "计算机软件著作权登记文档"
    doc.save(out_docx)


def write_markdown_docs(source_total: int, sample_lines: int, file_count: int) -> None:
    readme = """# 软件著作权申请材料包

本目录为 **星译漫画小说智能翻译软件 V1.1.0** 的软著申请准备材料。

## 先填三项

1. 申请人：个人姓名 / 公司全称（替换所有“【申请人名称，待填】”）。
2. 软件全称：建议“星译漫画小说智能翻译软件”，如商标检索有冲突再改名。
3. 开发完成日期：当前材料按 2026-10-31 编制，按实际填写。

## 文件清单

| 文件 | 用途 |
|---|---|
| 00-申请信息表.md | 在线填表字段草稿 |
| 01-软件说明书.docx | 文档鉴别材料，导出 PDF 后上传 |
| 02-源程序鉴别材料.docx | 源程序前后各 30 页，导出 PDF 后上传 |
| 02-源程序鉴别材料.txt | 源程序纯文本，便于核对与转 PDF |
| 03-第三方组件与模型清单.md | 开源/模型许可与原创范围说明 |
| 04-独立开发与权属声明.md | 独立开发/权属声明模板，签字后使用 |
| 05-提交材料清单与流程.md | 提交顺序、格式与检查项 |
| 06-命名与版本核对.md | 名称、版本、页眉一致性 |
| 07-合规与风险整改清单.md | 软著之外的上架/网信合规事项 |

## 已生成的源程序样本

- 纳入源文件：__FILE_COUNT__ 个（排除 build、oh_modules、llama.cpp、third_party、字体等第三方/生成物）。
- 原始有效行：约 __SOURCE_TOTAL__ 行。
- 样本：__SAMPLE_LINES__ 行，按每页 50 行共 60 页。
- 页眉：软件名称 + 版本；页脚：页码/总页数。
- 不含 API Key、密码或模型二进制。

## 重要提醒

- 本包是**申请准备材料**，不是登记机构出具的受理/证书。
- 源程序样本只用于登记鉴别，不能替代真实源代码仓库。
- 第三方组件与模型不纳入原创范围；申请前请核对许可。
- 申请人必须对“独立开发、权属清晰、无侵权”声明负责，签字前请逐项核实。
- DOCX 需在 Word/Pages 中导出 PDF；导出后核对页数、页眉和字体未错位。

## 重新生成

使用带 python-docx 的 Python：

~~~bash
/Users/kong/.dsh/dsh-runtimes/dsh-next-primary-runtime/dependencies/python/bin/python3 \
  softcopyright/tools/generate_softcopyright.py
~~~

修改脚本顶部的 SOFTWARE_NAME / VERSION / APPLICANT 后重新生成即可。
""".replace("__FILE_COUNT__", str(file_count)).replace("__SOURCE_TOTAL__", str(source_total)).replace("__SAMPLE_LINES__", str(sample_lines))
    (OUT / "README.md").write_text(readme, encoding="utf-8")

    info = """# 软件著作权登记申请信息表（草稿）

> 在线系统字段名可能随版本调整；以下内容用于提前准备，最终以中国版权保护中心系统为准。

## 一、软件基本信息

| 字段 | 建议填写 | 说明 |
|---|---|---|
| 软件全称 | 星译漫画小说智能翻译软件 | 与说明书、源程序页眉一致；如商标冲突再改 |
| 软件简称 | 星译 | 可选 |
| 版本号 | V1.1.0 | 与 AppScope/app.json5 的 versionName 对应 |
| 软件分类 | 应用软件 / 图像处理与翻译工具 | 以系统选项为准 |
| 开发方式 | 独立开发 | 如合作/委托/职务开发，另附权属证明 |
| 权利取得方式 | 原始取得 | 受让/继承需附合同或证明 |
| 开发完成日期 | 2026-10-31（按实际改） | 必须早于或等于申请日 |
| 首次发表日期 | 未发表（如已公开上架则填实际日期） | 仅在公开发表后填写 |
| 发表地点 | 中国 | 未发表可不填 |
| 著作权人 | 【申请人名称，待填】 | 个人身份证姓名 / 公司营业执照全称 |
| 证件类型/号码 | 【待填】 | 个人身份证 / 统一社会信用代码 |
| 联系人/电话/邮箱 | 【待填】 | 受理通知使用 |
| 地址/邮编 | 【待填】 | 证书邮寄/联系 |

## 二、开发环境

| 字段 | 内容 |
|---|---|
| 硬件环境 | HarmonyOS NEXT 手机/平板/2in1，RAM ≥4GB；本地模型建议 ≥8GB |
| 软件环境 | HarmonyOS NEXT API 26+、DevEco Studio 6、ArkTS/ArkUI、C++/CMake |
| 编程语言 | ArkTS、C++ |
| 源程序量 | 约 __SOURCE_TOTAL__ 行（以实际统计为准） |
| 开发工具 | DevEco Studio、hvigor、CMake、llama.cpp、ONNX Runtime |

## 三、开发目的与功能

**开发目的：**

为 HarmonyOS NEXT 用户提供一个本地优先的漫画/小说翻译工作台，将图片导入、文字识别、翻译、版式覆盖、校对和导出整合到一个应用中，减少在多工具之间来回切换。

**面向行业：**

移动应用 / 数字阅读 / 图像文字处理 / 机器翻译工具。

**主要功能：**

1. 本地漫画、小说导入与书架管理；
2. 漫画页 OCR（系统 OCR / 本地 PP-OCRv6）；
3. 在线翻译服务与本地 llama.cpp 翻译；
4. 按气泡/段落合并翻译并按页面坐标排版；
5. 逐块文字颜色、描边、自动字号与横竖排；
6. 长按气泡人工校对并保存；
7. 术语表与翻译记忆；
8. 原文/译文/双语 ZIP 与译文 PDF 导出；
9. 手机/平板/2in1 断点布局与沉浸光感界面。

**技术特点：**

- Stage 模型 + 多 HAR 模块；
- 原生 ONNX Runtime det+rec 桥，TaskPool 异步，主线程只解码；
- llama.cpp GGUF 本地推理；
- Canvas 页面坐标系排版；
- 断点驱动的多端布局。

## 四、权属与材料

- 申请人类型：个人 / 公司（待定）
- 是否职务开发：是 / 否（如“是”，需单位出具权属说明）
- 是否委托/合作开发：是 / 否（如“是”，需合同）
- 是否使用第三方开源：是（见 03 清单）
- 是否包含第三方模型：不随安装包分发；用户自行下载/导入

## 五、待确认

- [ ] 申请人全称与证件信息
- [ ] 软件名称商标检索
- [ ] 开发完成日期
- [ ] 是否已公开发表
- [ ] 上游 StarFlow 是否存在代码/素材复制
- [ ] 第三方模型许可是否允许当前分发方式
""".replace("__SOURCE_TOTAL__", str(source_total))
    (OUT / "00-申请信息表.md").write_text(info, encoding="utf-8")

    third = """# 第三方组件、模型与原创范围清单

> 本清单用于软著申请时区分“申请人原创代码”与“第三方组件”。  
> 软著登记不审查开源许可，但申请人不能把第三方代码/模型声称为自己原创；上架/商用前还需单独核对许可。

## 一、随安装包分发的第三方组件

| 组件 | 位置 | 许可证 | 用途 | 义务/备注 |
|---|---|---|---|---|
| llama.cpp | entry/src/main/cpp/llama.cpp | MIT | 本地 GGUF 推理 | 保留版权与许可声明 |
| ONNX Runtime | @ohos/onnxruntime / 原生 liblibrary.so | MIT | 本地 OCR 推理 | 保留版权与许可声明；本项目通过 C API 链接 |
| Noto Sans CJK SC | entry/src/main/resources/rawfile/fonts/NotoSansSC.otf | SIL OFL 1.1 | 覆盖层/导出字体 | 保留 OFL 文本；不得单独售卖字体 |
| HarmonyOS SDK / ArkUI / Core Vision Kit | 系统 | 华为 SDK 许可 | 系统能力 | 不随仓库分发源码 |
| @ohos/aki | oh_modules | 以其包内 LICENSE 为准 | NAPI 辅助 | 申请前核对 |

## 二、不随安装包分发的模型

| 模型 | 获取方式 | 用途 | 许可状态 |
|---|---|---|---|
| PP-OCRv6 det/rec | 应用内下载 | 本地 OCR | 当前下载地址指向上游 StarFlow 仓库；需核实模型许可与再分发条件 |
| Hy-MT2 GGUF | 用户导入/下载 | 本地翻译 | 以模型发布方许可为准；商用需单独确认 |

**软著申请口径：**

- 模型文件不作为源程序提交；
- 说明书中描述“用户自行下载/导入”，不声称模型为申请人原创；
- 若申请材料包含模型加载/推理代码，该部分代码属于软件作品，模型权重不属于。

## 三、上游 StarFlow

- 上游：https://github.com/xujunjiex/StarFlow （LGPL-3.0）
- 本项目声明为 HarmonyOS 独立实现，但功能信息架构与交互概念参照上游。
- 申请前必须逐项确认：
  - [ ] 是否复制了上游 Java/Kotlin 代码；
  - [ ] 是否复制了上游图片、图标、文案；
  - [ ] 是否直接使用上游模型文件；
  - [ ] 是否保留上游版权/许可声明。
- 如存在复制：取得许可或删除后重新实现，不能作为“纯原创”申请。

## 四、AI 辅助开发说明

本项目开发过程中使用了 AI 编程助手。软著要求人的智力创作，建议保留：

- 人工需求拆分、架构设计、接口定义；
- 人工代码审查、修改、测试记录；
- Git 提交历史与版本迭代记录；
- 关键模块的设计说明。

申请材料不应把 AI 单独生成、未经人工选择/修改的内容声称为独立创作。

## 五、申请人原创范围建议

建议在申请表中按以下口径描述：

> 申请人独立完成了本软件的 HarmonyOS 应用层、业务服务层、页面排版与交互、原生 OCR 桥接、构建配置及配套文档；第三方开源组件与模型按各自许可证使用，不计入申请人原创范围。
"""
    (OUT / "03-第三方组件与模型清单.md").write_text(third, encoding="utf-8")

    decl = """# 独立开发与权属声明（模板）

> 本文件是模板，必须由申请人根据真实情况修改、签字/盖章后使用。  
> 不要在未核实的情况下直接提交虚假声明。

## 声明人

- 姓名/单位：____________________
- 证件类型/号码：____________________
- 联系地址：____________________

## 声明事项

本人/本单位就“星译漫画小说智能翻译软件 V1.1.0”（以下简称“本软件”）声明如下：

1. 本软件由声明人组织/独立开发，开发过程真实，源代码、文档和构建产物来源于本项目仓库。
2. 本软件未通过反编译、破解、复制等方式使用他人软件代码；除已列明的第三方开源组件和模型外，未将他人代码或素材作为自己的原创成果申请登记。
3. 本软件使用的第三方组件、字体、模型均按各自许可证使用，相关声明见 NOTICE.md 与 03-第三方组件与模型清单.md。
4. 本软件不包含申请人无权使用的漫画、小说、图片等内容；用户导入内容由用户自行负责。
5. 如本软件属于职务开发、委托开发或合作开发，声明人已取得相应权属证明/授权，并将在申请材料中提交。
6. 声明人保证申请材料真实、准确；如因材料不实产生法律责任，由声明人承担。

## 签字/盖章

声明人（签字/盖章）：____________________

日期：____年__月__日

## 附件建议

- [ ] 个人身份证 / 公司营业执照
- [ ] 职务开发权属说明（如适用）
- [ ] 委托/合作开发合同（如适用）
- [ ] 第三方组件许可文本
- [ ] 上游代码/素材比对说明（如适用）
"""
    (OUT / "04-独立开发与权属声明.md").write_text(decl, encoding="utf-8")

    checklist = """# 提交材料清单与流程

## 一、提交前必须完成的核对

- [ ] 申请人全称、证件号码、地址、电话、邮箱准确。
- [ ] 软件全称、版本号在申请表、说明书封面、源程序页眉三处一致。
- [ ] 开发完成日期早于申请日；未发表则不填首次发表日期。
- [ ] 源程序为申请人原创部分；第三方代码/生成代码已剔除或说明。
- [ ] 源程序前后各 30 页、每页 50 行；本包已生成 60 页 DOCX。
- [ ] 说明书页眉、页码、截图清晰；导出 PDF 后检查无错位。
- [ ] 文档中不包含 API Key、密码、内网地址、个人信息。
- [ ] 申请人签字/盖章的独立开发与权属声明。
- [ ] 个人身份证/公司营业执照扫描件。
- [ ] 如职务/委托/合作开发，附权属证明。
- [ ] 软件名称已做商标/应用名称检索（至少检索第 9、42 类）。

## 二、在线申请流程

1. 进入中国版权保护中心软件著作权登记系统，注册/登录。
2. 在线填写申请表，保存申请号。
3. 打印申请确认签章页，签字/盖章后扫描上传。
4. 上传：
   - 源程序鉴别材料 PDF（本包 02 文件导出）；
   - 文档鉴别材料 PDF（本包 01 文件导出）；
   - 身份证明；
   - 权属证明/声明（如适用）。
5. 提交后等待受理；如收到补正通知，按补正意见修改后重新上传。
6. 取得证书后核对：软件名称、版本、著作权人、登记号、日期。

> 具体系统字段、格式、受理时限以中国版权保护中心当期要求为准；本包不承诺审批结果。

## 三、DOCX 转 PDF

1. 用 Microsoft Word 或 macOS Pages 打开 DOCX；
2. 另存为/导出 PDF；
3. 打开 PDF 检查：
   - 源程序是否正好 60 页；
   - 每页是否 50 行；
   - 页眉是否显示软件名称/版本；
   - 中文是否乱码、表格是否截断；
4. 若页数不对，优先调整字号/行距/页边距，不要删除代码行。

## 四、文件命名建议

~~~text
星译漫画小说智能翻译软件V1.1.0-源程序鉴别材料.pdf
星译漫画小说智能翻译软件V1.1.0-软件说明书.pdf
星译漫画小说智能翻译软件V1.1.0-申请确认签章页.pdf
~~~

## 五、提交后保存

- 保留申请号、受理通知、补正通知、证书 PDF；
- 保留本包生成时的 Git commit；
- 保留源程序完整仓库和构建产物，备查。
"""
    (OUT / "05-提交材料清单与流程.md").write_text(checklist, encoding="utf-8")

    naming = """# 命名与版本核对

## 当前事实（来自仓库）

- AppScope/app.json5：
  - bundleName：com.moe.starflow
  - versionName：1.1.0
  - versionCode：1010000
- 桌面名称：星译 StarFlow
- 仓库许可证：LGPL-3.0
- 主要语言：ArkTS、C++

## 建议申请名称

- 软件全称：**星译漫画小说智能翻译软件**
- 软件简称：**星译**
- 版本号：**V1.1.0**

## 需要统一的位置

- [ ] 申请表“软件全称”
- [ ] 说明书封面/页眉
- [ ] 源程序页眉
- [ ] 应用桌面名称/关于页
- [ ] 导出 PDF 文件名
- [ ] 软著证书领取后对外宣传名称

## 商标/名称风险

“星译 / StarFlow”与上游项目同名。申请软著不要求商标注册，但上架/商用前应：

1. 在中国商标网检索“星译”“StarFlow”在第 9 类（软件）、第 42 类（软件服务）的注册情况；
2. 检索应用商店是否已有同名应用；
3. 如冲突，准备备用名并同步修改申请材料；
4. 不要使用与上游相同/近似的图标和视觉标识。

## 版本号规则

建议与 AppScope 的 versionName 保持一致，使用 V1.1.0；后续升级软著版本时，重新申请或办理变更/补充登记，按登记机构要求操作。
"""
    (OUT / "06-命名与版本核对.md").write_text(naming, encoding="utf-8")

    risk = """# 软著之外的风险整改清单

> 软著登记不审查内容合规、APP 备案、算法备案；但应用上架/公开运营前需要单独处理。

## P0：影响软著申请

- [ ] 确认申请人/权属（个人、公司、职务/委托/合作）。
- [ ] 核对上游 StarFlow 是否存在代码/素材复制；如有，取得许可或重写。
- [ ] 核对 PP-OCRv6 / Hy-MT2 模型许可，停止使用来源不明的下载地址。
- [ ] 补全 NOTICE 与第三方组件许可文本。
- [ ] 清理源程序样本中的密钥、密码、个人信息。
- [ ] 统一软件名称、版本、页眉。

## P1：公开上架前（网信/工信/商店）

- [ ] APP 备案（工信部系统）。
- [ ] 隐私政策 + 首次启动同意；列明在线翻译服务商与数据出境情况。
- [ ] 用户协议：仅限有权使用的内容、禁止违法内容、AI 结果需校对。
- [ ] AI 生成内容标识：阅读页角标、PDF/ZIP 导出页脚/元数据。
- [ ] 向属地网信办咨询算法/深度合成/生成式 AI 备案与安全评估。
- [ ] 内容投诉/版权下架机制、未成年人保护/年龄分级。
- [ ] 确认应用商店类目与资质要求（阅读/工具/出版）。
- [ ] API Key 本地保护，不写日志。

## P2：后续

- [ ] 商标注册/名称调整。
- [ ] 真机全分辨率回归。
- [ ] 模型许可商用确认。
- [ ] 开源 LGPL 源码提供方式（如对外分发二进制）。
"""
    (OUT / "07-合规与风险整改清单.md").write_text(risk, encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    files = collect_sources()
    lines, total = build_source_lines(files)
    sample_lines = len(lines)
    (OUT / "source-stats.json").write_text(json.dumps({
        "software": SOFTWARE_NAME,
        "version": VERSION,
        "source_files": len(files),
        "total_lines_before_trim": total,
        "sample_lines": sample_lines,
        "sample_pages": SAMPLE_PAGES,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    write_source_docx(lines, OUT / "02-源程序鉴别材料.docx", OUT / "02-源程序鉴别材料.txt")
    write_manual(OUT / "01-软件说明书.docx")
    write_markdown_docs(total, sample_lines, len(files))
    print("generated", OUT)
    print("source files", len(files), "total", total, "sample", sample_lines)


if __name__ == "__main__":
    main()
